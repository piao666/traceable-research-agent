"""Regression coverage for the post-R9 repository consistency repairs."""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.agent.file_access_policy import DOCS_ROOT
from app.agent.planner import plan_task
from app.database import Base
from app.mcp.server import SKILL_RUNNER_SPEC, _exposed_tool_specs, _run_skill_workflow
from app.skills.registry import init_skill_registry
from app.tools.defaults import register_default_tools
from app.tools.file_reader import read_file
from app.trace import store
from app.trace import models as trace_models  # noqa: F401
from app.evidence import models as evidence_models  # noqa: F401
from app.memory import models as memory_models  # noqa: F401
from app.improvement import models as improvement_models  # noqa: F401


ROOT = Path(__file__).resolve().parents[1]


class SkillRunnerRepairTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        register_default_tools()
        init_skill_registry(ROOT / "workspace" / "skills")

    def test_skill_parameters_override_query_and_integer_default(self) -> None:
        plan = plan_task(
            "outer query",
            allowed_tools=["tavily_search", "report_writer"],
            planner_mode="deterministic",
            skill_name="quick_search",
            skill_parameters={"query": "parameter query", "max_results": 2},
        )

        self.assertEqual(plan["skill_parameters"], {"query": "parameter query", "max_results": 2})
        self.assertEqual(plan["steps"][0]["arguments"]["query"], "parameter query")
        self.assertEqual(plan["steps"][0]["arguments"]["max_results"], 2)

    def test_skill_parameters_reject_unknown_names_and_wrong_types(self) -> None:
        with self.assertRaisesRegex(ValueError, "Unknown Skill parameters"):
            plan_task(
                "query",
                planner_mode="deterministic",
                skill_name="quick_search",
                skill_parameters={"unknown": True},
            )
        with self.assertRaisesRegex(ValueError, "must be integer"):
            plan_task(
                "query",
                planner_mode="deterministic",
                skill_name="quick_search",
                skill_parameters={"max_results": "two"},
            )

    def test_skill_runner_discloses_local_side_effects(self) -> None:
        self.assertFalse(SKILL_RUNNER_SPEC.read_only)
        self.assertFalse(SKILL_RUNNER_SPEC.side_effect_free)
        runner = next(tool for tool in _exposed_tool_specs() if tool.name == "skill_runner")
        self.assertFalse(runner.read_only)
        self.assertFalse(runner.side_effect_free)
        self.assertTrue(runner.policy["mcp_exposable"])

    def test_skill_runner_passes_parameters_into_persisted_plan(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            engine = create_engine(f"sqlite:///{Path(temporary) / 'runner.sqlite'}")
            Base.metadata.create_all(engine)
            local_session = sessionmaker(autocommit=False, autoflush=False, bind=engine)
            captured: dict[str, object] = {}

            def fake_run_plan(db, run_id: str) -> dict[str, str]:
                run = store.get_agent_run(db, run_id)
                captured.update(json.loads(run.plan_json or "{}"))
                return {"status": "completed"}

            with (
                patch("app.database.SessionLocal", local_session),
                patch("app.agent.executor.run_plan", side_effect=fake_run_plan),
            ):
                result = _run_skill_workflow({
                    "skill_name": "quick_search",
                    "query": "outer query",
                    "parameters": {"query": "inner query", "max_results": 3},
                })

            engine.dispose()

        self.assertTrue(result.success, result.error_message)
        self.assertEqual(captured["skill_parameters"], {"query": "inner query", "max_results": 3})
        self.assertEqual(captured["steps"][0]["arguments"]["query"], "inner query")
        self.assertEqual(captured["steps"][0]["arguments"]["max_results"], 3)

    def test_skill_runner_rejects_non_object_parameters(self) -> None:
        result = _run_skill_workflow({
            "skill_name": "quick_search",
            "query": "query",
            "parameters": [],
        })
        self.assertFalse(result.success)
        self.assertEqual(result.metadata["error_type"], "invalid_args")

    def test_skill_runner_reports_parameter_validation_as_invalid_args(self) -> None:
        result = _run_skill_workflow({
            "skill_name": "quick_search",
            "query": "query",
            "parameters": {"max_results": "two"},
        })
        self.assertFalse(result.success)
        self.assertEqual(result.metadata["error_type"], "invalid_args")
        self.assertIn("must be integer", result.error_message)


class OfficeFileReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        DOCS_ROOT.mkdir(parents=True, exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=DOCS_ROOT)
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_docx_paragraphs_and_tables_are_extracted(self) -> None:
        path = self.root / "sample.docx"
        document = Document()
        document.add_paragraph("Traceable DOCX evidence")
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Name"
        table.rows[0].cells[1].text = "Score"
        table.rows[1].cells[0].text = "Alpha"
        table.rows[1].cells[1].text = "9"
        document.save(path)

        result = read_file({"path": str(path), "max_chars": 2000})

        self.assertTrue(result.success, result.error_message)
        self.assertIn("Traceable DOCX evidence", result.output["content"])
        self.assertEqual(result.output["tables"][0]["columns"], ["Name", "Score"])
        self.assertEqual(result.output["tables"][0]["rows"], [["Alpha", "9"]])
        self.assertEqual(result.metadata["extraction_method"], "docx")

    def test_xlsx_sheets_are_extracted_without_formula_execution(self) -> None:
        path = self.root / "sample.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Metrics"
        sheet.append(["Name", "Score"])
        sheet.append(["Alpha", 9])
        sheet.append(["Formula", "=1+1"])
        workbook.save(path)
        workbook.close()

        result = read_file({"path": str(path), "max_chars": 2000})

        self.assertTrue(result.success, result.error_message)
        self.assertIn("[Sheet: Metrics]", result.output["content"])
        self.assertEqual(result.output["tables"][0]["columns"], ["Name", "Score"])
        self.assertEqual(result.output["tables"][0]["rows"][0], ["Alpha", "9"])
        self.assertEqual(result.output["tables"][0]["rows"][1], ["Formula", ""])
        self.assertEqual(result.metadata["extraction_method"], "xlsx")

    def test_pdf_routes_to_dedicated_reader(self) -> None:
        path = self.root / "sample.pdf"
        path.write_bytes(b"%PDF-1.4\n")

        result = read_file({"path": str(path)})

        self.assertFalse(result.success)
        self.assertEqual(result.metadata["error_type"], "wrong_tool")
        self.assertIn("pdf_reader", result.error_message)


if __name__ == "__main__":
    unittest.main()
