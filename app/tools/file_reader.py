"""Safe local file reader tool."""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from app.agent.file_access_policy import (
    DOCS_ROOT,
    allowed_roots,
    display_path,
    find_allowed_root,
    resolve_file_reader_path,
)
from app.tools.base import ToolResult


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MAX_CHARS = 8000
MAX_CHARS_LIMIT = 20000
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".py", ".log"}
OFFICE_EXTENSIONS = {".docx", ".xlsx"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | OFFICE_EXTENSIONS
MAX_OFFICE_FILE_BYTES = 25 * 1024 * 1024
MAX_OFFICE_ARCHIVE_BYTES = 100 * 1024 * 1024
MAX_OFFICE_ARCHIVE_ENTRIES = 5000
MAX_OFFICE_TABLES = 20
MAX_TABLE_ROWS = 1000
MAX_TABLE_COLUMNS = 32
MAX_TABLE_CHARS = 50000
MAX_CELL_CHARS = 500


def _failure(
    message: str,
    *,
    error_type: str,
    path: str | None = None,
    extension: str | None = None,
    resolved_path: Path | None = None,
) -> ToolResult:
    return ToolResult(
        success=False,
        error_message=message,
        metadata={
            "error_type": error_type,
            "path": path,
            "extension": extension,
            "docs_root": str(DOCS_ROOT),
            "allowed_roots": [str(root) for root in allowed_roots()],
            "resolved_path_summary": str(resolved_path) if resolved_path else None,
        },
    )


def _coerce_max_chars(value: Any) -> int:
    try:
        max_chars = int(value) if value is not None else DEFAULT_MAX_CHARS
    except (TypeError, ValueError):
        max_chars = DEFAULT_MAX_CHARS
    return max(1, min(max_chars, MAX_CHARS_LIMIT))


def _output_path(resolved_path: Path, allowed_root: Path | None) -> str:
    if allowed_root is not None:
        try:
            return resolved_path.relative_to(allowed_root).as_posix()
        except ValueError:
            pass
    try:
        return resolved_path.relative_to(DOCS_ROOT).as_posix()
    except ValueError:
        return display_path(resolved_path)


def _approved_path_value(arguments: dict[str, Any]) -> str | None:
    value = arguments.get("_approved_file_reader_path")
    return str(value).strip() if value else None


def _validate_office_archive(path: Path) -> None:
    """Reject oversized or malformed OOXML archives before parsing them."""

    if path.stat().st_size > MAX_OFFICE_FILE_BYTES:
        raise ValueError(f"Office file exceeds {MAX_OFFICE_FILE_BYTES} bytes.")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_OFFICE_ARCHIVE_ENTRIES:
                raise ValueError("Office archive contains too many entries.")
            if sum(member.file_size for member in members) > MAX_OFFICE_ARCHIVE_BYTES:
                raise ValueError("Office archive expands beyond the safety limit.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Office file is not a valid OOXML archive.") from exc


def _bounded_lines(lines: Any, max_chars: int) -> tuple[str, bool]:
    output: list[str] = []
    length = 0
    for raw_line in lines:
        line = str(raw_line or "").strip()
        if not line:
            continue
        separator = 1 if output else 0
        available = max_chars - length - separator
        if available <= 0:
            return "\n".join(output), True
        if len(line) > available:
            output.append(line[:available])
            return "\n".join(output), True
        output.append(line)
        length += separator + len(line)
    return "\n".join(output), False


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def _table_payload(
    raw_rows: list[list[str]],
    *,
    caption: str,
    extraction_method: str,
    truncated: bool,
) -> dict[str, Any] | None:
    rows = [list(row) for row in raw_rows if any(cell for cell in row)]
    if not rows:
        return None
    width = min(max(len(row) for row in rows), MAX_TABLE_COLUMNS)
    normalized = [(row + [""] * width)[:width] for row in rows]
    columns = [cell[:MAX_CELL_CHARS] for cell in normalized[0]]
    columns = [name or f"column_{index + 1}" for index, name in enumerate(columns)]
    data_rows = [
        [cell[:MAX_CELL_CHARS] for cell in row]
        for row in normalized[1:MAX_TABLE_ROWS + 1]
    ]
    return {
        "columns": columns,
        "rows": data_rows,
        "caption": caption[:MAX_CELL_CHARS],
        "truncated": truncated or len(normalized) > MAX_TABLE_ROWS + 1,
        "extraction_method": extraction_method,
        "source_bound": True,
    }


def _read_docx(path: Path, max_chars: int) -> tuple[str, bool, list[dict[str, Any]]]:
    from docx import Document

    _validate_office_archive(path)
    document = Document(path)
    tables: list[dict[str, Any]] = []
    table_lines: list[str] = []
    remaining_table_chars = MAX_TABLE_CHARS
    for index, table in enumerate(document.tables[:MAX_OFFICE_TABLES], start=1):
        raw_rows: list[list[str]] = []
        table_truncated = len(table.rows) > MAX_TABLE_ROWS + 1
        for row in table.rows[:MAX_TABLE_ROWS + 1]:
            values = [_cell_text(cell.text) for cell in row.cells[:MAX_TABLE_COLUMNS]]
            row_chars = sum(len(value) for value in values)
            if row_chars > remaining_table_chars:
                table_truncated = True
                break
            remaining_table_chars -= row_chars
            table_truncated = table_truncated or len(row.cells) > MAX_TABLE_COLUMNS
            raw_rows.append(values)
        payload = _table_payload(
            raw_rows,
            caption=f"Table {index}",
            extraction_method="docx_table",
            truncated=table_truncated,
        )
        if payload is not None:
            tables.append(payload)
            table_lines.append(f"[Table {index}]")
            table_lines.extend("\t".join(row) for row in raw_rows)

    content, content_truncated = _bounded_lines(
        (paragraph.text for paragraph in document.paragraphs),
        max_chars,
    )
    if len(content) < max_chars and table_lines:
        table_content, table_content_truncated = _bounded_lines(
            table_lines,
            max_chars - len(content) - (1 if content else 0),
        )
        if table_content:
            content = f"{content}\n{table_content}" if content else table_content
        content_truncated = content_truncated or table_content_truncated
    elif table_lines:
        content_truncated = True
    return content, content_truncated or len(document.tables) > MAX_OFFICE_TABLES, tables


def _read_xlsx(path: Path, max_chars: int) -> tuple[str, bool, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    _validate_office_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    tables: list[dict[str, Any]] = []
    content_lines: list[str] = []
    remaining_table_chars = MAX_TABLE_CHARS
    workbook_truncated = len(workbook.sheetnames) > MAX_OFFICE_TABLES
    try:
        for worksheet in workbook.worksheets[:MAX_OFFICE_TABLES]:
            raw_rows: list[list[str]] = []
            table_truncated = (
                int(worksheet.max_row or 0) > MAX_TABLE_ROWS + 1
                or int(worksheet.max_column or 0) > MAX_TABLE_COLUMNS
            )
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=MAX_TABLE_ROWS + 1,
                max_col=MAX_TABLE_COLUMNS,
                values_only=True,
            ):
                values = [_cell_text(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if not values:
                    continue
                row_chars = sum(len(value) for value in values)
                if row_chars > remaining_table_chars:
                    table_truncated = True
                    break
                remaining_table_chars -= row_chars
                raw_rows.append(values)
            payload = _table_payload(
                raw_rows,
                caption=worksheet.title,
                extraction_method="xlsx_sheet",
                truncated=table_truncated,
            )
            if payload is not None:
                tables.append(payload)
                content_lines.append(f"[Sheet: {worksheet.title}]")
                content_lines.extend("\t".join(row) for row in raw_rows)
        content, content_truncated = _bounded_lines(content_lines, max_chars)
        return content, content_truncated or workbook_truncated, tables
    finally:
        workbook.close()


def _resolve_allowed_path_for_read(
    raw_path: str,
    arguments: dict[str, Any],
) -> tuple[Path | None, Path | None, bool, ToolResult | None]:
    approved_path = _approved_path_value(arguments)
    resolved = resolve_file_reader_path(raw_path)
    allowed_root = find_allowed_root(resolved)
    if allowed_root is not None:
        return resolved, allowed_root, False, None
    approved = False
    try:
        approved = bool(approved_path) and Path(approved_path).resolve() == resolved
    except OSError:
        approved = False
    if approved:
        return resolved, None, True, None
    error_type = "approval_mismatch" if approved_path else "safety_rejected"
    message = (
        "Approved file path did not match the requested file_reader path."
        if approved_path
        else "Path is outside configured file_reader allowed roots and has not been approved for this run."
    )
    return None, None, False, _failure(
        message,
        error_type=error_type,
        path=raw_path,
        extension=resolved.suffix.lower(),
        resolved_path=resolved,
    )


def read_file(arguments: dict[str, Any]) -> ToolResult:
    """Read a supported text or Office file under an allowed local root."""

    raw_path = str(arguments.get("path") or "").strip()
    max_chars = _coerce_max_chars(arguments.get("max_chars"))

    if not raw_path:
        return _failure("Missing required argument: path.", error_type="invalid_args")

    resolved_path, allowed_root, approved, failure = _resolve_allowed_path_for_read(
        raw_path, arguments
    )
    if failure is not None:
        return failure
    assert resolved_path is not None

    extension = resolved_path.suffix.lower()
    if extension == ".pdf":
        return _failure(
            "PDF files are supported by pdf_reader; call pdf_reader with this path.",
            error_type="wrong_tool",
            path=raw_path,
            extension=extension,
            resolved_path=resolved_path,
        )
    if extension not in SUPPORTED_EXTENSIONS:
        return _failure(
            f"Unsupported file extension: {extension or '<none>'}.",
            error_type="unsupported_format",
            path=raw_path,
            extension=extension,
            resolved_path=resolved_path,
        )
    if not resolved_path.exists() or not resolved_path.is_file():
        return _failure(
            "File not found.",
            error_type="not_found",
            path=raw_path,
            extension=extension,
            resolved_path=resolved_path,
        )

    try:
        if extension == ".docx":
            content, truncated, tables = _read_docx(resolved_path, max_chars)
            extraction_method = "docx"
        elif extension == ".xlsx":
            content, truncated, tables = _read_xlsx(resolved_path, max_chars)
            extraction_method = "xlsx"
        else:
            text = resolved_path.read_text(encoding="utf-8")
            truncated = len(text) > max_chars
            content = text[:max_chars]
            from app.tools.structured_tables import csv_tables
            tables = csv_tables(content, truncated=truncated) if extension == ".csv" else []
            extraction_method = "text"
    except UnicodeDecodeError:
        try:
            text = resolved_path.read_text(encoding="utf-8-sig")
            truncated = len(text) > max_chars
            content = text[:max_chars]
            from app.tools.structured_tables import csv_tables
            tables = csv_tables(content, truncated=truncated) if extension == ".csv" else []
            extraction_method = "text"
        except Exception as exc:
            return _failure(
                f"Failed to decode file: {exc}",
                error_type="read_error",
                path=raw_path,
                extension=extension,
                resolved_path=resolved_path,
            )
    except Exception as exc:
        return _failure(
            f"Failed to read file: {exc}",
            error_type="read_error",
            path=raw_path,
            extension=extension,
            resolved_path=resolved_path,
        )

    relative_path = _output_path(resolved_path, allowed_root)
    chars_read = len(content)
    return ToolResult(
        success=True,
        output={
            "path": relative_path,
            "content": content,
            "chars_read": chars_read,
            "truncated": truncated,
            "tables": tables,
        },
        output_summary=(
            f"Read {relative_path}: {chars_read} chars"
            + (" (truncated)" if truncated else "")
        ),
        metadata={
            "error_type": None,
            "safe_path": True,
            "approved_outside_allowed_roots": approved,
            "extension": extension,
            "docs_root": str(DOCS_ROOT),
            "allowed_root": str(allowed_root) if allowed_root else None,
            "allowed_roots": [str(root) for root in allowed_roots()],
            "resolved_path_summary": str(resolved_path),
            "max_chars": max_chars,
            "extraction_method": extraction_method,
        },
    )
